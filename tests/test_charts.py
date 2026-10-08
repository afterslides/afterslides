from __future__ import annotations

import datetime as dt
import io
import re

import openpyxl
import pytest

from afterslides import InvalidArgumentError, Presentation, Series, UnsupportedError, XySeries
from conftest import as_python_pptx, reopen, zip_entries


def embedded_workbook(prs: Presentation, chart_name: str) -> openpyxl.Workbook:
    """Loads the workbook behind a chart from the saved file."""
    entries = zip_entries(prs)
    slide = prs.shape(chart_name).slide
    rels = entries[slide.part_name.lstrip("/").replace("slides/", "slides/_rels/") + ".rels"]
    chart_targets = re.findall(rb'Target="\.\./charts/(chart\d+\.xml)"', rels)
    shapes = [s for s in slide.shapes if s.kind == "chart"]
    chart_file = chart_targets[[s.name for s in shapes].index(chart_name)].decode()
    chart_rels = entries[f"ppt/charts/_rels/{chart_file}.rels"]
    target = re.search(rb'Target="\.\./embeddings/([^"]+)"', chart_rels)
    assert target is not None
    return openpyxl.load_workbook(io.BytesIO(entries[f"ppt/embeddings/{target.group(1).decode()}"]))


def test_read_chart(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    assert chart.types == ["barChart"]
    assert chart.categories == ["Q1", "Q2", "Q3", "Q4"]
    assert chart.series == [
        Series("2024", [10.5, 12.0, 9.75, 14.0]),
        Series("2025", [11.0, 13.5, 10.0, 15.25]),
    ]
    assert chart.title == "Revenue"


def test_replace_data_with_mapping(prs: Presentation) -> None:
    prs.shape("Revenue Chart").chart.replace_data(
        ["North", "South", "East"],
        {"Plan": [1, 2, 3], "Actual": [1.5, None, 2.5], "Forecast": [4, 5, 6]},
    )
    check = as_python_pptx(prs).slides[1].shapes[1].chart
    plot = check.plots[0]
    assert list(plot.categories) == ["North", "South", "East"]
    assert [s.name for s in plot.series] == ["Plan", "Actual", "Forecast"]
    assert plot.series[1].values == (1.5, None, 2.5)

    wb = embedded_workbook(prs, "Revenue Chart")
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == (None, "Plan", "Actual", "Forecast")
    assert rows[1] == ("North", 1, 1.5, 4)
    assert rows[2] == ("South", 2, None, 5)
    # Number formats of the template carry over into the workbook.
    assert wb.active["B2"].number_format == "#,##0.00"


def test_new_series_copy_template_formatting(prs: Presentation) -> None:
    prs.shape("Revenue Chart").chart.replace_data(["a"], {"s1": [1], "s2": [2], "s3": [3]})
    xml = zip_entries(prs)["ppt/charts/chart1.xml"].decode()
    idx = re.findall(r'<c:idx val="(\d+)"/>', xml)
    assert len(idx) == len(set(idx)) == 3


def test_fewer_series_and_categories(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    chart.replace_data(["only"], [Series("one", [42])])
    chart = reopen(prs).shape("Revenue Chart").chart
    assert chart.categories == ["only"]
    assert chart.series == [Series("one", [42.0])]


def test_pandas_like_dataframe(prs: Presentation) -> None:
    class Frame:
        """Just enough of a DataFrame for duck typing."""

        index = ["Jan", "Feb"]
        columns = ["Visitors", "Buyers"]

        def __getitem__(self, col: str) -> list[float]:
            return {"Visitors": [10.0, float("nan")], "Buyers": [1.0, 2.0]}[col]

    prs.shape("Trend").chart.replace_data(Frame())
    chart = reopen(prs).shape("Trend").chart
    assert chart.series == [Series("Visitors", [10.0, None]), Series("Buyers", [1.0, 2.0])]


def test_numeric_and_date_categories(prs: Presentation) -> None:
    chart = prs.shape("Trend").chart
    chart.replace_data([2024, 2025], {"v": [1, 2]})
    assert reopen(prs).shape("Trend").chart.categories == [2024.0, 2025.0]
    chart.replace_data([dt.date(2026, 1, 1), dt.date(2026, 2, 1)], {"v": [1, 2]})
    assert reopen(prs).shape("Trend").chart.categories == [46023.0, 46054.0]


def test_pie_keeps_percent_format(prs: Presentation) -> None:
    prs.shape("Share").chart.replace_data(["x", "y"], {"Share": [0.25, 0.75]})
    assert embedded_workbook(prs, "Share").active["B2"].number_format == "0%"


def test_values_must_fit_categories(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    with pytest.raises(InvalidArgumentError):
        chart.replace_data(["a"], {"s": [1, 2]})
    with pytest.raises(InvalidArgumentError):
        chart.replace_data(["a"], {"s": ["one"]})


def test_scatter(prs: Presentation) -> None:
    chart = prs.shape("Scatter").chart
    assert chart.is_xy
    chart.replace_xy_data([XySeries("A", [1, 2, 3], [3, 2, 1]), XySeries("B", [0], [5])])
    xy = reopen(prs).shape("Scatter").chart.xy_series
    assert [s.name for s in xy] == ["A", "B"]
    assert xy[0].y == [3.0, 2.0, 1.0]
    check = as_python_pptx(prs).slides[4].shapes[0].chart.plots[0].series
    assert check[1].values == (5.0,)
    with pytest.raises(UnsupportedError):
        chart.replace_data(["a"], {"s": [1]})


def test_bubble(prs: Presentation) -> None:
    chart = prs.shape("Bubbles").chart
    chart.replace_xy_data([XySeries("M", [1, 2], [3, 4], sizes=[10, 20])])
    assert reopen(prs).shape("Bubbles").chart.xy_series == [
        XySeries("M", [1.0, 2.0], [3.0, 4.0], [10.0, 20.0])
    ]
    with pytest.raises(InvalidArgumentError):
        chart.replace_xy_data([XySeries("M", [1], [3])])


def test_title(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    chart.title = "Umsatz 2026"
    assert reopen(prs).shape("Revenue Chart").chart.title == "Umsatz 2026"
    check = as_python_pptx(prs).slides[1].shapes[1].chart
    assert check.chart_title.text_frame.text == "Umsatz 2026"
