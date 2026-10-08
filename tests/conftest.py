from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Callable, Union

import pptx
import pytest

import afterslides
from invariants import check_deck

FIXTURES = Path(__file__).parent / "fixtures"
TEMPLATE = FIXTURES / "template.pptx"


@pytest.fixture
def template_path() -> Path:
    return TEMPLATE


@pytest.fixture
def prs() -> afterslides.Presentation:
    return afterslides.Presentation(TEMPLATE)


_DUMP_DIR = os.environ.get("AFTERSLIDES_DUMP_DIR")
_dump_counter = 0


def dump(data: bytes, label: str | None = None, source: Path | None = None) -> None:
    """With AFTERSLIDES_DUMP_DIR set, keeps a copy of a saved deck there so CI
    can run external validators over everything the tests produce.

    `source` is the input deck; it goes into manifest.tsv so validation can
    ignore problems the input already had.
    """
    global _dump_counter
    if not _DUMP_DIR:
        return
    if label is None:
        label = os.environ.get("PYTEST_CURRENT_TEST", "unknown").split(" ")[0]
    _dump_counter += 1
    name = re.sub(r"[^A-Za-z0-9_.-]+", "_", label)[-150:]
    out_dir = Path(_DUMP_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{os.getpid()}-{_dump_counter:04d}-{name}.pptx"
    out.write_bytes(data)
    with open(out_dir / "manifest.tsv", "a", encoding="utf-8") as manifest:
        manifest.write(f"{out.name}\t{source.resolve() if source else TEMPLATE}\n")


def saved(deck: afterslides.Presentation) -> bytes:
    """Saves the deck and checks the structural invariants of the result."""
    data = deck.to_bytes()
    check_deck(data)
    dump(data)
    return data


def reopen(deck: afterslides.Presentation) -> afterslides.Presentation:
    return afterslides.Presentation(saved(deck))


def as_python_pptx(deck: afterslides.Presentation) -> pptx.presentation.Presentation:
    """Reads our output with python-pptx, an independent implementation."""
    return pptx.Presentation(io.BytesIO(saved(deck)))


def zip_entries(deck: afterslides.Presentation) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(saved(deck))) as z:
        return {name: z.read(name) for name in z.namelist()}


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert_with_libreoffice(pptx_path: Path, out_dir: Path, fmt: str = "pdf") -> Path:
    """Converts a deck with LibreOffice; a failure means it could not read it."""
    binary = soffice()
    assert binary is not None
    subprocess.run(
        [
            binary,
            f"-env:UserInstallation=file://{out_dir / 'profile'}",
            "--headless",
            "--convert-to",
            fmt,
            "--outdir",
            str(out_dir),
            str(pptx_path),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    out = out_dir / f"{pptx_path.stem}.{fmt}"
    assert out.exists() and out.stat().st_size > 0, f"LibreOffice did not produce a {fmt}"
    return out


Edit = Union[Callable[[str], str], str, bytes, None]


def patched(edits: dict[str, Edit], renames: dict[str, str] | None = None) -> bytes:
    """The template with some zip entries edited, added, renamed or removed.

    A callable receives the entry's text and returns the new text; str/bytes
    replace (or add) the entry; None removes it. Renames are applied first.
    """
    renames = renames or {}
    out = io.BytesIO()
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        seen = set()
        for info in src.infolist():
            name = renames.get(info.filename, info.filename)
            data = src.read(info.filename)
            seen.add(name)
            if name in edits:
                edit = edits[name]
                if edit is None:
                    continue
                if callable(edit):
                    data = edit(data.decode("utf-8")).encode("utf-8")
                else:
                    data = edit.encode("utf-8") if isinstance(edit, str) else edit
            dst.writestr(name, data)
        for name, edit in edits.items():
            if name not in seen and isinstance(edit, (str, bytes)):
                dst.writestr(name, edit)
    return out.getvalue()


def replace_once(old: str, new: str) -> Callable[[str], str]:
    def edit(text: str) -> str:
        assert old in text, f"{old!r} not found"
        return text.replace(old, new, 1)

    return edit
