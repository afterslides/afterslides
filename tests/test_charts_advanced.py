"""Combo charts, multi-level categories, number formats and dates."""

from __future__ import annotations

import datetime as dt
import io
import re

import openpyxl
import pytest
from test_regressions import FILTERED_SERIES, combo_chart

from afterslides import InvalidArgumentError, Presentation, Series, UnsupportedError
from conftest import as_python_pptx, patched, reopen, replace_once, zip_entries


@pytest.fixture
def combo() -> Presentation:
    return Presentation(patched({"ppt/charts/chart1.xml": combo_chart}))


def workbook(prs: Presentation, chart_file: str = "chart1.xml") -> openpyxl.Workbook:
    entries = zip_entries(prs)
    rels = entries[f"ppt/charts/_rels/{chart_file}.rels"].decode()
    target = re.search(r'Target="\.\./embeddings/([^"]+)"', rels)
    assert target is not None
    return openpyxl.load_workbook(io.BytesIO(entries[f"ppt/embeddings/{target.group(1)}"]))


# ---- combo charts ---------------------------------------------------------


def test_combo_series_report_their_plot(combo: Presentation) -> None:
    chart = combo.shape("Revenue Chart").chart
    assert [(s.name, s.plot) for s in chart.series] == [("2024", 0), ("2025", 0), ("2024", 1)]


def test_combo_grows_each_plot_from_its_own_template(combo: Presentation) -> None:
    combo.shape("Revenue Chart").chart.replace_data(
        ["a", "b"],
        [
            Series("col 1", [1, 2], plot=0),
            Series("line 1", [5, 6], plot=1),
            Series("col 2", [3, 4], plot=0),
            Series("line 2", [7, 8], plot=1),
            Series("col 3", [9, 9], plot=0),
        ],
    )
    plots = as_python_pptx(combo).slides[1].shapes[1].chart.plots
    assert [[s.name for s in p.series] for p in plots] == [
        ["col 1", "col 2", "col 3"],
        ["line 1", "line 2"],
    ]
    # Legend order follows the order the series were passed in.
    series = reopen(combo).shape("Revenue Chart").chart.series
    assert [(s.name, s.plot) for s in series] == [
        ("col 1", 0),
        ("line 1", 1),
        ("col 2", 0),
        ("line 2", 1),
        ("col 3", 0),
    ]
    xml = zip_entries(combo)["ppt/charts/chart1.xml"].decode()
    line = xml[xml.index("<c:lineChart>") :]
    assert "<c:invertIfNegative" not in line  # line copies came from the line series


def test_combo_shrinks_a_plot(combo: Presentation) -> None:
    combo.shape("Revenue Chart").chart.replace_data(
        ["a"], [Series("col", [1], plot=0), Series("line", [2], plot=1)]
    )
    plots = as_python_pptx(combo).slides[1].shapes[1].chart.plots
    assert [len(p.series) for p in plots] == [1, 1]


def test_combo_refuses_to_empty_a_plot(combo: Presentation) -> None:
    with pytest.raises(UnsupportedError, match="without series"):
        combo.shape("Revenue Chart").chart.replace_data(["a"], [Series("col", [1], plot=0)])


def test_plot_must_exist_and_be_set_everywhere(combo: Presentation) -> None:
    chart = combo.shape("Revenue Chart").chart
    with pytest.raises(InvalidArgumentError):
        chart.replace_data(["a"], [Series("x", [1], plot=0), Series("y", [1], plot=7)])
    with pytest.raises(InvalidArgumentError):
        chart.replace_data(["a"], [Series("x", [1], plot=0), Series("y", [1])])


def test_plot_index_on_a_simple_chart(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    chart.replace_data(["a"], [Series("x", [1], plot=0), Series("y", [2], plot=0)])
    assert [s.name for s in reopen(prs).shape("Revenue Chart").chart.series] == ["x", "y"]
    with pytest.raises(InvalidArgumentError):
        chart.replace_data(["a"], [Series("x", [1], plot=1)])


# ---- multi-level categories -----------------------------------------------


LEVELS = [
    ("2025", "Q3"),
    ("2025", "Q4"),
    ("2026", "Q1"),
    ("2026", "Q2"),
]


def test_multi_level_categories_round_trip(prs: Presentation) -> None:
    chart = prs.shape("Revenue Chart").chart
    chart.replace_data(LEVELS, {"Revenue": [1, 2, 3, 4]})
    assert reopen(prs).shape("Revenue Chart").chart.categories == LEVELS
    # python-pptx reads the hierarchy the same way.
    cats = as_python_pptx(prs).slides[1].shapes[1].chart.plots[0].categories
    assert cats.depth == 2
    assert [c.label for c in cats.levels[1]] == ["2025", "2026"]
    assert list(cats) == ["Q3", "Q4", "Q1", "Q2"]


def test_multi_level_workbook_layout(prs: Presentation) -> None:
    prs.shape("Revenue Chart").chart.replace_data(LEVELS, {"Revenue": [1, 2, 3, 4]})
    rows = list(workbook(prs).active.iter_rows(values_only=True))
    assert rows == [
        (None, None, "Revenue"),
        ("2025", "Q3", 1),
        (None, "Q4", 2),
        ("2026", "Q1", 3),
        (None, "Q2", 4),
    ]
    xml = zip_entries(prs)["ppt/charts/chart1.xml"].decode()
    assert "<c:f>Sheet1!$A$2:$B$5</c:f>" in xml
    assert "<c:f>Sheet1!$C$2:$C$5</c:f>" in xml


def test_outer_label_repeats_when_inner_group_changes(prs: Presentation) -> None:
    # The same outer label in two separate runs starts two groups.
    levels = [("A", "x"), ("B", "x"), ("A", "y")]
    prs.shape("Revenue Chart").chart.replace_data(levels, {"s": [1, 2, 3]})
    assert reopen(prs).shape("Revenue Chart").chart.categories == levels


def test_multi_level_needs_consistent_depth(prs: Presentation) -> None:
    with pytest.raises(InvalidArgumentError):
        prs.shape("Revenue Chart").chart.replace_data([("a", "b"), ("c",)], {"s": [1, 2]})


def test_pandas_multiindex_like_frame(prs: Presentation) -> None:
    class Frame:
        index = LEVELS
        columns = ["Revenue"]

        def __getitem__(self, col: str) -> list[float]:
            return [1.0, 2.0, 3.0, 4.0]

    prs.shape("Revenue Chart").chart.replace_data(Frame())
    assert reopen(prs).shape("Revenue Chart").chart.categories == LEVELS


# ---- formats and dates ----------------------------------------------------


def test_number_format_override(prs: Presentation) -> None:
    prs.shape("Revenue Chart").chart.replace_data(
        ["a", "b"], [Series("share", [0.25, 0.5], number_format="0.0%")]
    )
    series = reopen(prs).shape("Revenue Chart").chart.series
    assert series[0].number_format == "0.0%"
    assert workbook(prs).active["B2"].number_format == "0.0%"


def test_template_number_format_is_reported(prs: Presentation) -> None:
    assert prs.shape("Revenue Chart").chart.series[0].number_format == "#,##0.00"


def test_dates_get_a_date_format(prs: Presentation) -> None:
    prs.shape("Trend").chart.replace_data(
        [dt.date(2026, 1, 1), dt.date(2026, 2, 1)], {"Visitors": [1, 2]}
    )
    xml = zip_entries(prs)["ppt/charts/chart2.xml"].decode()
    assert "<c:formatCode>yyyy\\-mm\\-dd</c:formatCode>" in xml
    assert workbook(prs, "chart2.xml").active["A2"].is_date


# ---- filtered series ------------------------------------------------------


def test_filtered_series_are_dropped_on_replace() -> None:
    prs = Presentation(
        patched({"ppt/charts/chart1.xml": replace_once("</c:barChart>", FILTERED_SERIES)})
    )
    prs.shape("Revenue Chart").chart.replace_data(["a"], {"x": [1]})
    xml = zip_entries(prs)["ppt/charts/chart1.xml"].decode()
    assert "filteredBarSeries" not in xml
    assert "<c:extLst>" not in xml[: xml.index("</c:barChart>")]
