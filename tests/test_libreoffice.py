"""Opens edited decks in LibreOffice. A conversion failure means another
office suite could not read what we wrote."""

from __future__ import annotations

from pathlib import Path

import pptx
import pytest

from afterslides import Presentation, Series, XySeries
from conftest import TEMPLATE, convert_with_libreoffice, soffice

pytestmark = [
    pytest.mark.libreoffice,
    pytest.mark.skipif(soffice() is None, reason="LibreOffice is not installed"),
]


def edit_everything(prs: Presentation) -> None:
    prs.replace_text({"{{title}}": "Report", "{{customer}}": "Acme", "{{period}}": "Q3"})
    prs.shape("Revenue Chart").chart.replace_data(
        ["A", "B", "C"], [Series("x", [1, 2, 3]), Series("y", [3, 2, 1]), Series("z", [2, 2, 2])]
    )
    prs.shape("Top Customers").table.fill([[f"c{i}", i, f"{i}%"] for i in range(8)])
    prs.shape("Scatter").chart.replace_xy_data([XySeries("s", [1, 2], [2, 1])])
    prs.shape("Remove Me").delete()
    prs.slides[1].duplicate()
    prs.slides[4].delete()


def test_libreoffice_renders_edited_deck(tmp_path: Path) -> None:
    prs = Presentation(TEMPLATE)
    edit_everything(prs)
    out = tmp_path / "edited.pptx"
    prs.save(out)
    pdf = convert_with_libreoffice(out, tmp_path)
    assert pdf.read_bytes().startswith(b"%PDF")


def test_libreoffice_reads_every_slide(tmp_path: Path) -> None:
    prs = Presentation(TEMPLATE)
    edit_everything(prs)
    src = tmp_path / "src" / "edited.pptx"
    src.parent.mkdir()
    prs.save(src)
    # LibreOffice re-saves the deck; python-pptx then counts what it understood.
    resaved = convert_with_libreoffice(src, tmp_path, "pptx")
    check = pptx.Presentation(str(resaved))
    assert len(check.slides) == 7
    texts = [s.text_frame.text for slide in check.slides for s in slide.shapes if s.has_text_frame]
    assert "Report for Acme — Q3" in texts
