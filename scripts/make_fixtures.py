"""Builds the test template decks in tests/fixtures/.

The decks are generated with python-pptx so they can be rebuilt from source.
They deliberately contain the awkward bits real templates have: placeholders
split across runs, combo-free charts of every common type, grouped shapes,
notes, an internal hyperlink and sections.

Run with:  uv run python scripts/make_fixtures.py
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures"

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def tiny_png() -> bytes:
    """A 2x2 red PNG, so we don't need image files in the repo."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def textbox(slide, name, text, left=1, top=1, width=4, height=1):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = name
    box.text_frame.text = text
    return box


def add_sections(prs, names_and_slides):
    """Adds PowerPoint 2010 sections (python-pptx has no API for them)."""
    pres = prs.part._element
    ext_lst = pres.find(f"{{{P}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, f"{{{P}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{P}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    section_lst = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for i, (name, slides) in enumerate(names_and_slides):
        section = etree.SubElement(
            section_lst,
            f"{{{P14}}}section",
            name=name,
            id=f"{{6C1D0A3B-0000-4000-8000-00000000000{i}}}",
        )
        lst = etree.SubElement(section, f"{{{P14}}}sldIdLst")
        for slide in slides:
            etree.SubElement(lst, f"{{{P14}}}sldId", id=str(slide.slide_id))


def build_template() -> Presentation:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    # 1: title slide with placeholders split across runs.
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "{{title}}"
    sub = s1.placeholders[1].text_frame.paragraphs[0]
    for text, bold in [("Report for {{cust", True), ("omer}}", False), (" — {{period}}", False)]:
        run = sub.add_run()
        run.text = text
        run.font.bold = bold

    # 2: clustered column chart + note text + speaker notes.
    s2 = prs.slides.add_slide(title_only)
    s2.shapes.title.text = "Revenue by quarter"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3", "Q4"]
    data.add_series("2024", (10.5, 12.0, 9.75, 14.0), number_format="#,##0.00")
    data.add_series("2025", (11.0, 13.5, 10.0, 15.25), number_format="#,##0.00")
    frame = s2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6), Inches(4.5), data
    )
    frame.name = "Revenue Chart"
    chart = frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue"
    chart.plots[0].series[0].format.fill.solid()
    textbox(s2, "Note", "Source: {{source}}", 6.6, 1.5, 3, 1)
    s2.notes_slide.notes_text_frame.text = "Speaker notes for {{title}}"

    # 3: table with a header row and two data rows.
    s3 = prs.slides.add_slide(title_only)
    s3.shapes.title.text = "Top customers"
    tbl_frame = s3.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(9), Inches(1.2))
    tbl_frame.name = "Top Customers"
    table = tbl_frame.table
    for c, header in enumerate(["Customer", "Revenue", "Share"]):
        table.cell(0, c).text = header
    for r, row in enumerate([["Acme", "1.2M", "30%"], ["Globex", "0.8M", "20%"]], start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    # 4: line + pie charts, and a shape that tests delete.
    s4 = prs.slides.add_slide(blank)
    line = CategoryChartData()
    line.categories = ["Jan", "Feb", "Mar"]
    line.add_series("Visitors", (100, 120, 90))
    s4.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(0.5), Inches(4.5), Inches(3), line
    ).name = "Trend"
    pie = CategoryChartData()
    pie.categories = ["A", "B", "C"]
    pie.add_series("Share", (0.5, 0.3, 0.2), number_format="0%")
    s4.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(5), Inches(0.5), Inches(4.5), Inches(3), pie
    ).name = "Share"
    textbox(s4, "Remove Me", "{{optional}}", 0.5, 4, 4, 1)

    # 5: scatter and bubble charts.
    s5 = prs.slides.add_slide(blank)
    xy = XyChartData()
    series = xy.add_series("Measurements")
    for x, y in [(1, 2.5), (2, 3.5), (3, 2.0)]:
        series.add_data_point(x, y)
    s5.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER, Inches(0.5), Inches(0.5), Inches(4.5), Inches(3), xy
    ).name = "Scatter"
    bubble = BubbleChartData()
    bs = bubble.add_series("Markets")
    for x, y, size in [(1, 1, 5), (2, 3, 10)]:
        bs.add_data_point(x, y, size)
    s5.shapes.add_chart(
        XL_CHART_TYPE.BUBBLE, Inches(5), Inches(0.5), Inches(4.5), Inches(3), bubble
    ).name = "Bubbles"

    # 6: group shape and a picture.
    s6 = prs.slides.add_slide(blank)
    group = s6.shapes.add_group_shape()
    group.name = "Card"
    inner = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    inner.name = "Grouped Label"
    inner.text_frame.text = "Hello {{name}}"
    pic = s6.shapes.add_picture(io.BytesIO(tiny_png()), Inches(5), Inches(1), Inches(1), Inches(1))
    pic.name = "Logo"
    pic._element.nvPicPr.cNvPr.set("descr", "logo:company")

    # 7: hyperlink to slide 4.
    s7 = prs.slides.add_slide(blank)
    link = textbox(s7, "Link", "Go to charts", 1, 1, 4, 1)
    link.click_action.target_slide = s4

    add_sections(prs, [("Intro", [s1, s2, s3]), ("Charts", [s4, s5]), ("Appendix", [s6, s7])])
    return prs


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_template().save(OUT / "template.pptx")
    print(f"wrote {OUT / 'template.pptx'}")


if __name__ == "__main__":
    main()
